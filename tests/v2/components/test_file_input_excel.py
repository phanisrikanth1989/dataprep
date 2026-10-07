"""Excel file input, against v1 on the same job config and workbook.

Workbooks are built here, with openpyxl, when the tests run. v1 reads a
sheet through pandas and reads nothing from it when pandas raises, whatever
the cause; the cases compared with v1 are ones pandas reads. What v2 does
with the rest is tested on v2 alone.
"""
import datetime as dt
import io
import json
import os
from pathlib import Path

import openpyxl
import pytest

from src.v2 import load_job, run_job
from src.v2.errors import JobRefusedError
from tests.v2.answer_key import REPO_ROOT, assert_matches_v1

from .kit import columns, flow, job, writer

ABSENT = ...  # as a config value: leave the key out
LEGACY = REPO_ROOT / "tests" / "fixtures" / "data" / "sample_legacy.xls"
SAMPLE = REPO_ROOT / "tests" / "talend_xml_samples" / "converted_jsons" / "Job_tFileInputExcel_0.1.json"

DAY = dt.datetime(2024, 1, 31)
MOMENT = dt.datetime(2024, 1, 31, 10, 11, 12)


def book(sheets):
    """The bytes of a workbook: ``{sheet name: rows}``, a row being a list of cell values."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    held = io.BytesIO()
    workbook.save(held)
    return held.getvalue()


def excel(schema, write_schema="same", reject=False, titles=False, type_name="FileInputExcel", **config):
    """workbook -> file, optionally with the reader's reject output written to rej.csv.

    ``titles`` has the writer put the column names first; ``header`` is the reader's own key.
    """
    made = {"filepath": "in.xlsx", "die_on_error": False}
    made.update(config)
    config = {key: value for key, value in made.items() if value is not ABSENT}
    source = {"id": "in", "type": type_name, "config": config,
              "schema": {"input": [], "output": columns(schema) if schema else []}, "inputs": [], "outputs": ["row1"]}
    components = [source, writer(schema if write_schema == "same" else write_schema, inputs=("row1",),
                                 include_header=titles)]
    flows = [flow("row1", "in", "out")]
    if reject:
        source["outputs"].append("bad")
        components.append(writer(None, component_id="rej", path="rej.csv", inputs=("bad",)))
        flows.append(flow("bad", "in", "rej", "reject"))
    return job(components, flows)


def files(sheets):
    if sheets is None:
        return {}
    if isinstance(sheets, dict) and all(isinstance(rows, list) for rows in sheets.values()):
        return {"in.xlsx": book(sheets)}
    return sheets


def same(tmp_path, sheets, schema, fails=False, **kwargs):
    """Both engines do the same with the job; ``fails`` says the job is one neither finishes."""
    run = assert_matches_v1(excel(schema, **kwargs), files(sheets), tmp_path)
    assert run.succeeded is not fails, run.error
    return run


def v2(tmp_path, sheets, schema, **kwargs):
    """Run on v2 only, inside tmp_path; returns (result, the folder)."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    for name, data in files(sheets).items():
        (tmp_path / name).write_bytes(data if isinstance(data, bytes) else Path(data).read_bytes())
    previous = os.getcwd()
    os.chdir(tmp_path)
    try:
        return run_job(excel(schema, **kwargs)), tmp_path
    finally:
        os.chdir(previous)


def written(tmp_path, sheets, schema, name="out.csv", **kwargs):
    """What v2 alone writes for the job."""
    result, folder = v2(tmp_path, sheets, schema, **kwargs)
    assert result.status == "success", result.error
    return (folder / name).read_bytes()


def refused(schema="a:str, n:int", **kwargs):
    with pytest.raises(JobRefusedError) as caught:
        load_job(excel(schema, **kwargs))
    return caught.value.report.format()


def labelled(cells):
    """One sheet of two columns: what the cell is, and the cell."""
    return {"S": [[label, cell] for label, cell in cells]}


# ------------------------------------------------------------------
# A cell as each declared type
# ------------------------------------------------------------------

ALL_TYPES = "id:int, name:str, amt:float, ok:bool, d:datetime@%Y-%m-%d, m:Decimal#2"
PEOPLE = [
    ["id", "name", "amt", "ok", "d", "m"], [1, "alice", 10.5, True, DAY, 12.345], [2, "bob", 20, False, MOMENT, 7],
]


def test_every_type_is_read(tmp_path):
    run = same(tmp_path, {"S": PEOPLE}, ALL_TYPES, header=1)
    assert run.files["out.csv"] == b"1;alice;10.5;true;2024-01-31;12.35\n2;bob;20.0;false;2024-01-31;7.00\n"


def test_empty_cells_of_every_type(tmp_path):
    run = same(tmp_path, {"S": PEOPLE + [[3, None, None, None, None, 0]]}, ALL_TYPES, header=1)
    assert run.files["out.csv"].splitlines()[2] == b"3;;;false;;0.00"


def test_empty_decimal_is_missing(tmp_path):
    sheets = {"S": PEOPLE + [[3, None, None, None, None, None]]}
    same(tmp_path, sheets, ALL_TYPES, header=1, write_schema=None)
    # With a writer that declares the column v1 writes the text <NA>.
    assert written(tmp_path / "declared", sheets, ALL_TYPES, header=1).splitlines()[2] == b"3;;;false;;"


TEXT_CELLS = [
    ("whole", 30), ("fraction", 1.5), ("negative", -7), ("zero", 0), ("negative fraction", -1.7), ("third", 1 / 3),
    ("long", 12345678.9), ("longer", 123456789.1), ("sum", 0.1 + 0.2), ("small", 0.00012345), ("smaller", 1e-5),
    ("smaller still", 0.0000123456789), ("tiny", 1.5e-7), ("tiniest", 1.2345e-10), ("negative tiny", -2e-6),
    ("big", 1e16), ("bigger", 1e21), ("two to the 53", 2 ** 53),
    ("text", "abc"), ("number as text", "30"), ("fraction as text", "30.0"), ("padded", " 12 "), ("blank", "  "),
    ("word true", "true"), ("stamp as text", "2024-01-31 10:11:12"), ("small as text", "0.00001"),
    ("true", True), ("false", False),
    ("day", DAY), ("moment", MOMENT), ("date", dt.date(2024, 1, 31)), ("first day", dt.datetime(1900, 1, 1)),
    ("time", dt.time(10, 30)), ("time with seconds", dt.time(23, 59, 58)), ("midnight", dt.time(0, 0)),
    ("nothing", None), ("error", "#N/A"), ("tab", "a\tb"),
]


def test_cells_of_every_kind_as_text(tmp_path):
    run = same(tmp_path, labelled(TEXT_CELLS), "label:str, v:str")
    wrote = dict(line.split(";", 1) for line in run.files["out.csv"].decode().splitlines())
    assert wrote["whole"] == "30" and wrote["third"] == "0.3333333333333333" and wrote["smaller"] == "1e-05"
    assert wrote["smaller still"] == "1.23456789e-05" and wrote["tiny"] == "1.5e-07"
    assert wrote["bigger"] == "1" + 21 * "0"
    assert wrote["true"] == "1" and wrote["false"] == "0" and wrote["word true"] == "true"
    assert wrote["moment"] == "31-01-2024" and wrote["time"] == "10:30:00"
    assert wrote["stamp as text"] == "2024-01-31 10:11:12"
    assert wrote["nothing"] == "" and wrote["error"] == "" and wrote["blank"] == "  " and wrote["padded"] == " 12 "


NUMBER_CELLS = [
    ("whole", 30), ("fraction", 1.5), ("negative fraction", -1.7), ("third", 1 / 3), ("long", 123456789.1),
    ("big", 1e16), ("two to the 53", 2 ** 53), ("zero", 0),
    ("number as text", "30"), ("fraction as text", "30.0"), ("padded", " 12.5 "), ("exponent", "1e3"), ("signed", "+5"),
    ("zeros", "007"), ("half", ".5"), ("dot", "5."), ("long as text", "9007199254740993"),
    ("text", "abc"), ("comma", "1,5"), ("word nan", "nan"), ("blank", "  "), ("nothing", None), ("error", "#DIV/0!"),
    ("true", True), ("false", False), ("day", DAY), ("time", dt.time(10, 30)),
]


@pytest.mark.parametrize("kind", ["int", "float", "float#2"])
def test_cells_of_every_kind_as_a_number(tmp_path, kind):
    run = same(tmp_path, labelled(NUMBER_CELLS), f"label:str, v:{kind}")
    wrote = dict(line.split(";", 1) for line in run.files["out.csv"].decode().splitlines())
    assert wrote["text"] == "" and wrote["true"] == "" and wrote["day"] == "" and wrote["comma"] == ""
    assert wrote["negative fraction"] == ("-1" if kind == "int" else "-1.7")


def test_infinity_as_a_float(tmp_path):
    same(tmp_path, labelled([("a", "inf"), ("b", "-Infinity"), ("c", 1.5)]), "label:str, v:float")


DECIMAL_CELLS = [cell for cell in NUMBER_CELLS if cell[0] not in ("word nan",)] + [
    ("round up", 2.675), ("round up too", 1.005), ("round down", -1.005), ("sum", 0.1 + 0.2), ("small", 1e-5),
    ("many places as text", "1234567890123456.789"),
]


@pytest.mark.parametrize("places", [0, 2, 4])
def test_cells_of_every_kind_as_a_decimal(tmp_path, places):
    run = same(tmp_path, labelled(DECIMAL_CELLS), f"label:str, v:Decimal#{places}", write_schema=None)
    wrote = dict(line.split(";", 1) for line in run.files["out.csv"].decode().splitlines())
    assert wrote["text"] == "" and wrote["true"] == "" and wrote["day"] == ""
    if places == 2:
        assert wrote["round up"] == "2.68" and wrote["round up too"] == "1.01" and wrote["round down"] == "-1.01"
        assert wrote["many places as text"] == "1234567890123456.79"


def test_decimal_without_declared_places(tmp_path):
    cells = [("a", 30), ("b", 1.5), ("c", 12345678.9), ("d", "30.0"), ("e", "1.50"), ("f", "007"), ("g", -0.25)]
    run = same(tmp_path, labelled(cells), "label:str, v:Decimal")
    assert run.files["out.csv"] == b"a;30\nb;1.5\nc;12345678.9\nd;30\ne;1.5\nf;7\ng;-0.25\n"


BOOL_CELLS = [
    ("true", True), ("false", False), ("one", 1), ("zero", 0), ("two", 2), ("fraction", 1.5),
    ("word true", "true"), ("upper", "TRUE"), ("padded yes", " Yes "), ("on", "on"), ("no", "no"), ("one as text", "1"),
    ("zero as text", "0"), ("text", "abc"), ("y", "y"), ("day", DAY), ("nothing", None), ("blank", "  "),
]


def test_cells_of_every_kind_as_a_bool(tmp_path):
    run = same(tmp_path, labelled(BOOL_CELLS), "label:str, v:bool")
    wrote = dict(line.split(";", 1) for line in run.files["out.csv"].decode().splitlines())
    yes = {label for label, value in wrote.items() if value == "true"}
    assert yes == {"true", "one", "word true", "upper", "padded yes", "on", "one as text"}


DATE_CELLS = [
    ("day", DAY), ("moment", MOMENT), ("date", dt.date(2024, 2, 29)),
    ("thousandths", dt.datetime(2024, 1, 31, 10, 11, 12, 345000)),
    ("first day", dt.datetime(1900, 1, 1)), ("march 1900", dt.datetime(1900, 3, 1)), ("far", dt.datetime(9999, 12, 31)),
    ("as text", "2024-01-31"), ("other pattern", "31/01/2024"), ("padded", " 2024-01-31 "),
    ("stamp as text", "2024-01-31 10:11:12"), ("text", "abc"), ("number", 20240131), ("fraction", 1.5), ("true", True),
    ("time", dt.time(10, 30)), ("nothing", None), ("blank", "  "), ("no such day", "2024-02-30"),
]


def test_cells_of_every_kind_as_a_date(tmp_path):
    run = same(tmp_path, labelled(DATE_CELLS), "label:str, v:datetime@%Y-%m-%d",
               write_schema="label:str, v:datetime@%Y-%m-%d %H:%M:%S.%f")
    wrote = dict(line.split(";", 1) for line in run.files["out.csv"].decode().splitlines())
    assert wrote["moment"] == "2024-01-31 10:11:12.000000" and wrote["thousandths"] == "2024-01-31 10:11:12.345000"
    assert wrote["as text"] == "2024-01-31 00:00:00.000000"
    unread = ("other pattern", "padded", "stamp as text", "text", "number", "true", "time", "nothing", "no such day")
    assert [label for label in unread if wrote[label] != ""] == []


@pytest.mark.parametrize(
    "pattern, cells",
    [
        ("%Y%m%d", [20240131, "20240228", None]),
        ("%d/%m/%Y %H:%M:%S", ["31/01/2024 10:11:12", MOMENT]),
        ("", [DAY, MOMENT, None]),
        ("", ["2024-01-31", "2024-02-29", None]),
        ("", ["2024-01-31 10:11:12", None]),
        ("", ["31/01/2024", "29/02/2024"]),
    ],
)
def test_date_patterns(tmp_path, pattern, cells):
    kind = f"datetime@{pattern}" if pattern else "datetime"
    same(tmp_path, {"S": [[cell, index] for index, cell in enumerate(cells)]}, f"d:{kind}, n:int",
         write_schema="d:datetime@%Y-%m-%d %H:%M:%S, n:int")


def test_date_written_with_the_writers_pattern(tmp_path):
    run = same(tmp_path, {"S": [[MOMENT, 1]]}, "d:datetime@%d/%m/%Y %H:%M:%S, n:int")
    assert run.files["out.csv"] == b"31/01/2024 10:11:12;1\n"


def test_each_date_is_read_by_the_declared_pattern_alone(tmp_path):
    # When no cell fits the pattern v1 tries others on the whole column; one cell changes how the rest are read.
    out = written(tmp_path, {"S": [["31/01/2024", 1], ["2024-02-29", 2]]}, "d:datetime@%Y-%m-%d, n:int")
    assert out == b";1\n2024-02-29;2\n"
    assert written(tmp_path / "none", {"S": [["31/01/2024", 1]]}, "d:datetime@%Y-%m-%d, n:int") == b";1\n"


def test_date_without_a_pattern_is_read_cell_by_cell(tmp_path):
    # v1 guesses one format for the column here and reads a number as a count of nanoseconds.
    cells = [["2024-01-31", 1], ["2024-02-29 10:00:00", 2], ["29/02/2024", 3], [45000, 4], [DAY, 5]]
    out = written(tmp_path, {"S": cells}, "d:datetime, n:int", write_schema="d:datetime@%Y-%m-%d %H:%M:%S, n:int")
    assert out == b"2024-01-31 00:00:00;1\n2024-02-29 10:00:00;2\n2024-02-29 00:00:00;3\n;4\n2024-01-31 00:00:00;5\n"


def test_date_type_is_read_like_datetime_and_cut_to_the_day(tmp_path):
    # The converter never writes this type; what v1 makes of it changes from cell to cell.
    cells = [[MOMENT, 1], ["2024-02-29", 2], [dt.time(10, 30), 3], ["abc", 4], [None, 5]]
    assert written(tmp_path, {"S": cells}, "d:date@%Y-%m-%d, n:int") == b"2024-01-31;1\n2024-02-29;2\n;3\n;4\n;5\n"


def test_duration_cells_are_read_as_dates(tmp_path):
    # v1 (openpyxl) knows a duration by its cell format: "1 day, 2:00:00" as text, missing as a date.
    # The views v2 reads a sheet under do not tell a duration from a date.
    cells = labelled([("long", dt.timedelta(days=1, hours=2)), ("short", dt.timedelta(hours=5, minutes=30))])
    assert written(tmp_path, cells, "label:str, v:str") == b"long;01-01-1900\nshort;05:30:00\n"


def test_formulas_are_read_by_their_last_result(tmp_path):
    xlsxwriter = pytest.importorskip("xlsxwriter")
    held = io.BytesIO()
    workbook = xlsxwriter.Workbook(held, {"in_memory": True})
    sheet = workbook.add_worksheet("S")
    sheet.write_formula(0, 0, "=1+1", None, 2)
    sheet.write_formula(0, 1, '="a"&"b"', None, "ab")
    sheet.write_formula(0, 2, "=1/3", None, 1 / 3)
    sheet.write_formula(1, 0, "=1/0", None, "#DIV/0!")
    sheet.write_formula(1, 1, "=TRUE()", None, True)
    sheet.write_formula(1, 2, "=2.5*2", None, 5)
    workbook.close()
    run = same(tmp_path, {"in.xlsx": held.getvalue()}, "n:int, s:str, f:float")
    assert run.files["out.csv"] == b"2;ab;0.3333333333333333\n;1;5.0\n"


def test_int_column_stays_whole_whatever_else_it_holds(tmp_path):
    # In v1 one number too large for the type leaves the whole column as floats: 30 is written 30.0.
    assert written(tmp_path, labelled([("a", 30), ("b", 1e21), ("c", 7)]), "label:str, v:int") == b"a;30\nb;\nc;7\n"
    # And the text inf in the column has v1 read no row of the sheet at all.
    assert written(tmp_path / "inf", labelled([("a", 30), ("b", "inf")]), "label:str, v:int") == b"a;30\nb;\n"


# ------------------------------------------------------------------
# Which rows are read
# ------------------------------------------------------------------

TEXT = "a:str, b:str, c:str"
GRID = [[f"{letter}{number}" for letter in "abcde"] for number in (1, 2, 3, 4)]


@pytest.mark.parametrize(
    "config",
    [
        {},
        {"header": 1},
        {"header": "1"},
        {"header": True},
        {"header": False},
        {"header": 2},
        {"header": 4},
        {"header": 9},
        {"footer": 1},
        {"header": 1, "footer": 2},
        {"footer": 4},
        {"footer": 9},
        {"limit": "2"},
        {"limit": 2},
        {"limit": "0"},
        {"limit": 0},
        {"limit": ""},
        {"limit": None},
        {"limit": "9"},
        {"header": 1, "limit": "2"},
    ],
)
def test_header_footer_and_limit(tmp_path, config):
    same(tmp_path, {"S": GRID}, TEXT, **config)


def test_limit_and_footer_together(tmp_path):
    # v1 reads nothing here: pandas refuses the two together and v1 swallows the refusal.
    assert written(tmp_path, {"S": GRID}, TEXT, header=1, footer=1, limit="1") == b"a2;b2;c2\n"
    assert written(tmp_path / "wide", {"S": GRID}, TEXT, footer=1, limit="9") == b"a1;b1;c1\na2;b2;c2\na3;b3;c3\n"


def test_footer_given_as_text(tmp_path):
    # v1 accepts the text and then reads nothing: pandas wants a number.
    assert written(tmp_path, {"S": GRID}, TEXT, footer="3") == b"a1;b1;c1\n"


@pytest.mark.parametrize(
    "config",
    [{"header": -1}, {"header": "x"}, {"footer": -1}, {"footer": "x"}, {"limit": -1}, {"limit": "abc"},
     {"first_column": -1}, {"first_column": "B"}],
)
def test_counts_that_cannot_be_fail_on_both(tmp_path, config):
    same(tmp_path, {"S": GRID}, TEXT, fails=True, **config)
    assert next(iter(config)) in refused(TEXT, **config)


def test_header_and_first_column_can_come_from_the_context(tmp_path):
    made = excel(TEXT, header="${context.skip}", first_column="context.start")
    made["context"] = {"Default": {"skip": {"value": "1", "type": "int"}, "start": {"value": "2", "type": "str"}}}
    run = assert_matches_v1(made, files({"S": GRID}), tmp_path)
    assert run.files["out.csv"] == b"b2;c2;d2\nb3;c3;d3\nb4;c4;d4\n"


def test_footer_and_limit_can_come_from_the_context(tmp_path):
    # v1 checks these two before it looks up the context, and fails on the reference.
    made = excel(TEXT, footer="${context.foot}", limit="context.most")
    made["context"] = {"Default": {"foot": {"value": "1", "type": "int"}, "most": {"value": "2", "type": "str"}}}
    (tmp_path / "in.xlsx").write_bytes(book({"S": GRID}))
    previous = os.getcwd()
    os.chdir(tmp_path)
    try:
        result = run_job(made)
    finally:
        os.chdir(previous)
    assert result.status == "success", result.error
    assert (tmp_path / "out.csv").read_bytes() == b"a1;b1;c1\na2;b2;c2\n"


def test_empty_rows_are_rows_except_at_the_end(tmp_path):
    empty = [None, None, None]
    rows = [empty, empty, ["a1", "b1", "c1"], empty, ["a2", "b2", "c2"], empty, empty]
    run = same(tmp_path, {"S": rows}, TEXT)
    assert run.files["out.csv"] == b";;\n;;\na1;b1;c1\n;;\na2;b2;c2\n"
    for name, config in (("header", {"header": 2}), ("footer", {"footer": 1}), ("limit", {"limit": "3"})):
        same(tmp_path / name, {"S": rows}, TEXT, **config)


@pytest.mark.parametrize("header", [0, 1, 2])
@pytest.mark.parametrize("limit", ["1", "2", "3", "4", "5", "9"])
def test_limit_stops_short_of_the_empty_rows_it_ends_on(tmp_path, header, limit):
    empty = [None, None, None]
    filled = ["a", "b", "c"]
    rows = [filled, empty, filled, empty, empty, filled, empty, [None, None, None, "d8"]]
    same(tmp_path, {"S": rows}, TEXT, header=header, limit=limit)


def test_row_that_holds_something_beyond_the_columns_read_is_not_empty(tmp_path):
    filled, beyond, empty = ["a1", "b1"], [None, None, None, None, "e"], [None, None]
    run = same(tmp_path, {"S": [filled, beyond, beyond, filled]}, "a:str, b:str", limit="2")
    assert run.files["out.csv"] == b"a1;b1\n;\n"
    run = same(tmp_path / "bare", {"S": [filled, empty, empty, filled]}, "a:str, b:str", limit="2")
    assert run.files["out.csv"] == b"a1;b1\n"


def test_limit_on_a_sheet_of_empty_rows_only_at_the_top(tmp_path):
    sheets = {"S": [[None, None], [None, None], ["a3", "b3"]]}
    assert same(tmp_path, sheets, "a:str, b:str", limit="1").files["out.csv"] == b""
    assert same(tmp_path / "two", sheets, "a:str, b:str", limit="2").files["out.csv"] == b";\n;\n"


def test_a_row_of_blanks_at_the_end_is_a_row(tmp_path):
    same(tmp_path, {"S": [["a1", "b1", "c1"], ["  ", None, None]]}, TEXT)


@pytest.mark.parametrize("config", [{}, {"header": 2}, {"limit": "2"}, {"footer": 1}])
def test_sheet_of_one_column(tmp_path, config):
    same(tmp_path, {"S": [["a1"], [None], ["a2"], ["  "], ["a3"]]}, "a:str", **config)


def test_stopread_on_emptyrow_is_ignored(tmp_path):
    run = same(tmp_path, {"S": [["a", 1], [None, None], ["b", 2]]}, "a:str, n:int", stopread_on_emptyrow=True)
    assert run.files["out.csv"] == b"a;1\n;\nb;2\n"


# ------------------------------------------------------------------
# Which columns are read
# ------------------------------------------------------------------

@pytest.mark.parametrize(
    "config",
    [
        {"first_column": 1},
        {"first_column": 2},
        {"first_column": "2"},
        {"first_column": 3},
        {"last_column": "3"},
        {"last_column": 3},
        {"last_column": "C"},
        {"last_column": "c"},
        {"first_column": 2, "last_column": "D"},
        {"first_column": 2, "last_column": "E"},
        {"first_column": 3, "last_column": "AA"},
        {"last_column": "10"},
        {"last_column": ""},
        {"last_column": None},
    ],
)
def test_first_and_last_column(tmp_path, config):
    same(tmp_path, {"S": GRID}, TEXT, **config)


def test_columns_are_counted_from_the_first_of_the_sheet(tmp_path):
    run = same(tmp_path, {"S": [[None, "x1", "y1"], [None, "x2", "y2"]]}, TEXT)
    assert run.files["out.csv"] == b";x1;y1\n;x2;y2\n"
    shifted = {"S": [[None, None, "x1", "y1", "z1"], [None, None, "x2", "y2", "z2"]]}
    same(tmp_path / "third", shifted, TEXT, first_column=3)
    same(tmp_path / "left", {"S": [[None, None, None, "d1"], [None, None, None, "d2"]]}, TEXT)


def test_rows_of_different_lengths(tmp_path):
    run = same(tmp_path, {"S": [["a1", "b1", "c1"], ["a2"], ["a3", "b3"], ["a4", "b4", "c4", "d4"]]}, TEXT)
    assert run.files["out.csv"] == b"a1;b1;c1\na2;;\na3;b3;\na4;b4;c4\n"


@pytest.mark.parametrize(
    "rows, config, expected",
    [
        ([["a1", "b1"], ["a2", "b2"]], {}, b"a1;b1;\na2;b2;\n"),
        (GRID[:2], {"last_column": "2"}, b"a1;b1;\na2;b2;\n"),
        (GRID[:2], {"first_column": 4}, b"d1;e1;\nd2;e2;\n"),
        (GRID[:2], {"first_column": 6}, b";;\n;;\n"),
    ],
)
def test_columns_the_sheet_or_the_range_does_not_hold_are_missing(tmp_path, rows, config, expected):
    # v1 reads no row at all: pandas refuses a range that is not the size of the schema, and v1 swallows that.
    assert written(tmp_path, {"S": rows}, TEXT, **config) == expected


def test_first_column_zero_is_the_first_column(tmp_path):
    # v1 refuses 0, which its own converter sample holds.
    assert written(tmp_path, {"S": GRID[:1]}, TEXT, first_column=0) == b"a1;b1;c1\n"


@pytest.mark.parametrize(
    "config", [{"first_column": 3, "last_column": "2"}, {"last_column": "A1"}, {"last_column": "-1"}]
)
def test_last_column_that_makes_no_range_is_refused(config):
    # v1 ignores such a last column and reads from the first column of the sheet.
    assert "last_column" in refused(TEXT, **config)


# ------------------------------------------------------------------
# Which sheets are read
# ------------------------------------------------------------------

PAIR = "a:str, n:int"


def named(name):
    """A ``sheetlist`` entry that names a sheet."""
    return {"sheetname": name}


def like(pattern):
    """A ``sheetlist`` entry that is a regular expression."""
    return {"sheetname": pattern, "use_regex": True}

BOOK = {"Q1": [["x1", 1], ["x2", 2]], "Q2": [["y1", 3]], "Other": [["z1", 4]], "q3 data": [["w1", 5]]}


@pytest.mark.parametrize(
    "config, expected",
    [
        ({}, b"x1;1\nx2;2\n"),
        ({"all_sheets": True}, b"x1;1\nx2;2\ny1;3\nz1;4\nw1;5\n"),
        ({"sheetlist": [{"sheetname": "Q2", "use_regex": False}]}, b"y1;3\n"),
        ({"sheetlist": [named("Q2"), named("Other")]}, b"y1;3\n"),
        ({"sheetlist": ["Other"]}, b"z1;4\n"),
        ({"sheetlist": [named("Nope")]}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [named("Nope")], "die_on_error": True}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [named("q2")]}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [named("")]}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [named("1")]}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [like("Q.*")]}, b"x1;1\nx2;2\n"),
        ({"sheetlist": [like("2$")]}, b"y1;3\n"),
        ({"sheetlist": [like("^zz")]}, b""),
        ({"all_sheets": True, "sheetlist": [named("Other")]}, b"z1;4\n"),
        ({"all_sheets": True, "sheetlist": [named("oth")]}, b"z1;4\n"),
        ({"all_sheets": True, "sheetlist": [named("Q1")]}, b"x1;1\nx2;2\n"),
        ({"all_sheets": True, "sheetlist": [like("th")]}, b"z1;4\n"),
        ({"all_sheets": True, "sheetlist": [named("Q2"), named("Q2"), like("2$")]}, b"y1;3\n"),
        ({"all_sheets": True, "sheetlist": [named("zzz")]}, b""),
        ({"all_sheets": True, "sheetlist": ["Q2", "Nope"]}, b"y1;3\n"),
    ],
)
def test_sheet_selection(tmp_path, config, expected):
    assert same(tmp_path, BOOK, PAIR, **config).files["out.csv"] == expected


@pytest.mark.parametrize(
    "sheetlist, expected",
    [
        ([named("Other"), named("Q1")], b"x1;1\nx2;2\nz1;4\n"),
        ([named("q")], b"x1;1\nx2;2\ny1;3\nw1;5\n"),
        ([like("^Q"), named("data")], b"x1;1\nx2;2\ny1;3\nw1;5\n"),
    ],
)
def test_chosen_sheets_are_read_in_workbook_order(tmp_path, sheetlist, expected):
    # v1 reads them in the order of a Python set, which changes from run to run.
    assert written(tmp_path, BOOK, PAIR, all_sheets=True, sheetlist=sheetlist) == expected


@pytest.mark.parametrize("config", [{"sheetlist": [like("^zz")]}, {"all_sheets": True, "sheetlist": [named("zzz")]}])
def test_no_sheet_to_read_fails_when_errors_are_fatal(tmp_path, config):
    same(tmp_path, BOOK, PAIR, die_on_error=True, fails=True, **config)
    result, _ = v2(tmp_path / "direct", BOOK, PAIR, die_on_error=True, **config)
    assert result.error == "No sheets found to read"


def test_sheet_pattern_that_is_no_pattern_is_refused():
    # v1 reads no sheet.
    assert "sheetlist" in refused(sheetlist=[like("(")])
    load_job(excel("a:str", sheetlist=[{"sheetname": "(", "use_regex": False}]))


def test_sheet_name_can_come_from_the_context(tmp_path):
    made = excel(PAIR, sheetlist=[{"sheetname": "${context.sheet}", "use_regex": False}])
    made["context"] = {"Default": {"sheet": {"value": "Other", "type": "str"}}}
    run = assert_matches_v1(made, files(BOOK), tmp_path)
    assert run.files["out.csv"] == b"z1;4\n"


@pytest.mark.parametrize(
    "config", [{"header": 1}, {"limit": "1"}, {"footer": 1}, {"first_column": 2, "last_column": "B"}]
)
def test_rows_and_columns_are_chosen_sheet_by_sheet(tmp_path, config):
    schema = "n:int" if "first_column" in config else PAIR
    same(tmp_path, BOOK, schema, all_sheets=True, **config)


def test_sheets_with_nothing_in_them_are_skipped(tmp_path):
    run = same(tmp_path, {"A": [["x1", 1]], "E": [], "B": [["y1", 2, "more"]]}, PAIR, all_sheets=True)
    assert run.files["out.csv"] == b"x1;1\ny1;2\n"


def test_hidden_sheets_are_read_and_chart_sheets_hold_no_rows(tmp_path):
    from openpyxl.chart import BarChart, Reference

    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = "Data"
    for row in (["a", 1], ["b", 2]):
        data.append(row)
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=2))
    workbook.create_chartsheet("Chart").add_chart(chart)
    for name, state in (("Hidden", "hidden"), ("Very", "veryHidden")):
        sheet = workbook.create_sheet(name)
        sheet.append([name, 9])
        sheet.sheet_state = state
    held = io.BytesIO()
    workbook.save(held)
    run = same(tmp_path, {"in.xlsx": held.getvalue()}, PAIR, all_sheets=True)
    assert run.files["out.csv"] == b"a;1\nb;2\nHidden;9\nVery;9\n"


def test_name_of_the_last_sheet_read_is_kept(tmp_path):
    result, _ = v2(tmp_path, BOOK, PAIR, all_sheets=True, sheetlist=[like("Q")])
    assert result.global_map["in_CURRENT_SHEET"] == "Q2"


# ------------------------------------------------------------------
# Files
# ------------------------------------------------------------------

@pytest.mark.parametrize("die_on_error", [False, ABSENT])
@pytest.mark.parametrize("sheets", [None, {"in.xlsx": b"not a workbook"}])
def test_file_that_is_missing_or_no_workbook_reads_no_rows(tmp_path, sheets, die_on_error):
    run = same(tmp_path, sheets, PAIR, die_on_error=die_on_error, titles=True)
    assert run.files["out.csv"] == b"a;n\n"


def test_missing_file_fails_the_reader_when_errors_are_fatal(tmp_path):
    same(tmp_path, None, PAIR, die_on_error=True, fails=True)
    result, folder = v2(tmp_path / "direct", None, PAIR, die_on_error=True)
    assert result.status == "failed" and result.failed_component == "in"
    assert result.error == "Excel file not found: in.xlsx"
    assert not (folder / "out.csv").exists()


def test_file_that_is_no_workbook_fails_the_reader_when_errors_are_fatal(tmp_path):
    same(tmp_path, {"in.xlsx": b"not a workbook"}, PAIR, die_on_error=True, fails=True)
    result, _ = v2(tmp_path / "direct", {"in.xlsx": b"not a workbook"}, PAIR, die_on_error=True)
    assert result.failed_component == "in" and result.error.startswith("Error reading Excel file in.xlsx: ")


def test_missing_file_with_a_reject_output_writes_two_empty_files(tmp_path):
    result, folder = v2(tmp_path, None, PAIR, reject=True, titles=True)
    assert result.status == "success", result.error
    assert (folder / "out.csv").read_bytes() == b"a;n\n"
    assert (folder / "rej.csv").read_bytes() == b"a;n;errorCode;errorMessage\n"


@pytest.mark.parametrize("path", ['"in.xlsx"', " 'in.xlsx' ", " in.xlsx "])
def test_quotes_and_blanks_around_the_path_are_dropped(tmp_path, path):
    assert same(tmp_path, BOOK, PAIR, filepath=path).files["out.csv"] == b"x1;1\nx2;2\n"


@pytest.mark.parametrize("name", ["in.xlsm", "in.XLSX"])
def test_other_file_names(tmp_path, name):
    same(tmp_path, {name: book(BOOK)}, PAIR, filepath=name)


def test_path_can_come_from_the_context(tmp_path):
    made = excel(PAIR, filepath="context.excelFile")
    made["context"] = {"Default": {"excelFile": {"value": "data/sales.xlsx", "type": "str"}}}
    assert_matches_v1(made, {"data/sales.xlsx": book(BOOK)}, tmp_path)


@pytest.mark.parametrize(
    "schema, config, expected",
    [
        ("id:int, name:str", {"header": 1}, b"1;Old1\n2;Old2\n3;Old3\n"),
        ("id:str, name:str", {}, b"id;name\n1;Old1\n2;Old2\n3;Old3\n"),
        ("id:float, name:str", {"header": 1, "footer": 1}, b"1.0;Old1\n2.0;Old2\n"),
        ("id:Decimal#2, name:str", {"header": 2, "limit": "1"}, b"2.00;Old2\n"),
        ("id:int, name:str", {"header": 1, "limit": "0"}, b""),
        ("id:int, name:str", {"limit": "2"}, b";name\n1;Old1\n"),
        ("id:int, name:str", {"header": 3, "limit": "5"}, b"3;Old3\n"),
        ("id:bool, name:str", {"header": 1, "sheetlist": [named("Nope")]}, b"true;Old1\nfalse;Old2\nfalse;Old3\n"),
        ("name:str", {"header": 1, "first_column": 2, "all_sheets": True, "sheetlist": [named("leg")]},
         b"Old1\nOld2\nOld3\n"),
    ],
)
def test_xls_workbook(tmp_path, schema, config, expected):
    run = same(tmp_path, {"in.xls": LEGACY}, schema, filepath="in.xls", **config)
    assert run.files["out.csv"] == expected


def test_large_sheet_keeps_every_row_in_order(tmp_path):
    rows = [[number, f"name {number}", number / 4, number % 2 == 0, DAY] for number in range(3000)]
    run = same(tmp_path, {"S": rows}, "n:int, s:str, f:float, b:bool, d:datetime@%Y-%m-%d", header=1, footer=1)
    lines = run.files["out.csv"].splitlines()
    assert len(lines) == 2998 and lines[0] == b"1;name 1;0.25;false;2024-01-31"
    assert lines[-1] == b"2998;name 2998;749.5;true;2024-01-31"


# ------------------------------------------------------------------
# Values that may not be missing
# ------------------------------------------------------------------

NEEDED = "a:str!, f:float!"
GAPS = {"S": [["a", 1.5], ["b", None], [None, 2.5], ["d", "x"]]}


def test_missing_value_where_none_is_allowed_is_rejected(tmp_path):
    run = same(tmp_path, GAPS, NEEDED, reject=True, titles=True)
    assert run.files["out.csv"] == b"a;f\na;1.5\n;2.5\n"
    assert run.files["rej.csv"] == (
        b"a;f;errorCode;errorMessage\n"
        b"b;;SCHEMA_VIOLATION;Column 'f': non-nullable column has null\n"
        b"d;;SCHEMA_VIOLATION;Column 'f': non-nullable column has null\n"
    )


def test_missing_value_where_none_is_allowed_fails_when_errors_are_fatal(tmp_path):
    same(tmp_path, GAPS, NEEDED, die_on_error=True, fails=True)
    result, folder = v2(tmp_path / "direct", GAPS, NEEDED, die_on_error=True)
    assert result.failed_component == "in"
    assert result.error == "Column 'f' has NULL values but is not nullable; the row is row 2 of sheet 'S' of in.xlsx"
    assert not (folder / "out.csv").exists()


def test_missing_value_where_none_is_allowed_is_rejected_when_the_key_is_left_out(tmp_path):
    # v1 fails the job here: its base class takes errors as fatal when the key is absent.
    out = written(tmp_path, GAPS, NEEDED, die_on_error=ABSENT, reject=True, name="rej.csv")
    assert out.splitlines()[1] == b"b;;SCHEMA_VIOLATION;Column 'f': non-nullable column has null"


def test_missing_date_where_none_is_allowed_is_rejected(tmp_path):
    # v1 lets an empty date through: it is still text when its base class looks for missing values.
    out = written(tmp_path, {"S": [[DAY, 1], [None, 2]]}, "d:datetime!@%Y-%m-%d, n:int", reject=True, name="rej.csv")
    assert out.splitlines()[1] == b";2;SCHEMA_VIOLATION;Column 'd': non-nullable column has null"


def test_whole_numbers_stay_whole_when_rows_are_rejected(tmp_path):
    # v1 writes the int column of the rows it keeps as floats (1.0) once a row is rejected.
    out = written(tmp_path, {"S": [[1, 1.5], [None, 2.5], [3, 3.5]]}, "n:int!, f:float", reject=True)
    assert out == b"1;1.5\n3;3.5\n"


def test_missing_bool_where_none_is_allowed_is_rejected(tmp_path):
    run = same(tmp_path, {"S": [["a", True], ["b", None], ["c", "no"]]}, "a:str, b:bool!", reject=True)
    assert run.files["out.csv"] == b"a;true\nc;false\n"
    assert run.files["rej.csv"].splitlines()[1] == b"b;;SCHEMA_VIOLATION;Column 'b': non-nullable column has null"


def test_clean_sheet_with_a_reject_output_writes_an_empty_reject_file(tmp_path):
    # v1 stalls here (status "error") because the reject flow got nothing.
    for name, schema in (("needed", NEEDED), ("free", "a:str, f:float")):
        result, folder = v2(tmp_path / name, {"S": [["a", 1.5]]}, schema, reject=True, titles=True)
        assert result.status == "success", result.error
        assert (folder / "rej.csv").read_bytes() == b"a;f;errorCode;errorMessage\n"
        assert (folder / "out.csv").read_bytes() == b"a;f\na;1.5\n"


# ------------------------------------------------------------------
# What v2 says no to, and what it lets pass
# ------------------------------------------------------------------

def test_keys_v1_never_acts_on_are_accepted(tmp_path):
    cells = {"S": [[" x1 ", "1.234,5", DAY], ["x2", "7", "2024-01-31"]]}
    run = same(
        tmp_path, cells, "a:str, n:str, d:str",
        trimall=True, trim_select=[{"column": "a", "trim": True}],
        advanced_separator=True, thousands_separator=".", decimal_separator=",",
        convertdatetostring=True, date_select=[{"column": "d", "convert_date": True, "pattern": "yyyy-MM-dd"}],
        password="secret", version_2007=True, affect_each_sheet="x", novalidate_on_cell=True, suppress_warn=True,
        encoding="UTF-8", read_real_value=True, generation_mode="EVENT_MODE", include_phoneticruns=False,
        configure_inflation_ratio=True, inflation_ratio="0.01",
    )
    assert run.files["out.csv"] == b" x1 ;1.234,5;31-01-2024\nx2;7;2024-01-31\n"


def test_reader_without_a_schema_is_refused():
    assert "schema" in refused(schema=None, write_schema=None)


def test_unknown_key_is_refused():
    assert "bogus" in refused(bogus=1)


@pytest.mark.parametrize("type_name", ["FileInputExcel", "tFileInputExcel", "file_input_excel"])
def test_type_names(tmp_path, type_name):
    assert written(tmp_path, BOOK, PAIR, type_name=type_name) == b"x1;1\nx2;2\n"


def test_v2_spelling_of_the_path():
    made = excel(PAIR)
    made["components"][0]["config"]["path"] = made["components"][0]["config"].pop("filepath")
    load_job(made)


def converters_sample():
    """The job of the converter's own sample, its log component swapped for a file output."""
    sample = json.loads(SAMPLE.read_text(encoding="utf-8"))
    source = next(component for component in sample["components"] if component["type"] == "FileInputExcel")
    target = writer(None, inputs=("row1",))
    target["schema"]["input"] = source["schema"]["output"]
    return job([source, target], [flow("row1", source["id"], "out")], context=sample["context"])


def test_the_converters_sample_loads():
    load_job(converters_sample())


def test_the_converters_sample_runs(tmp_path):
    # v1 refuses its first_column of 0. Its last_column of 5 leaves the sixth declared column without a source.
    rows = [["sale_id", "product", "region", "quantity", "revenue", "sale_date"], [7, "pen", "north", 3, 4.5, DAY]]
    target = tmp_path / "data" / "sales_data.xlsx"
    target.parent.mkdir()
    target.write_bytes(book({"Sales": rows}))
    previous = os.getcwd()
    os.chdir(tmp_path)
    try:
        result = run_job(converters_sample())
    finally:
        os.chdir(previous)
    assert result.status == "success", result.error
    assert (tmp_path / "out.csv").read_bytes() == (
        b"sale_id;product;region;quantity;revenue;sale_date\n7;pen;north;3;4.5;\n"
    )
