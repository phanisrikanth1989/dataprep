"""Excel output, against v1 on the same job config: what is in the workbook each engine writes."""
import datetime as dt
import io
import os
import zipfile

import openpyxl
import pytest

from src.v2 import load_job, run_job
from src.v2.errors import JobRefusedError
from tests.v2 import answer_key

from .kit import columns, flow, job, reader
from .test_file_input_excel import book

SCHEMA = "id:int, name:str, amt:float"
DATA = b"id;name;amt\n1;ann;1.5\n2;bob;2\n3;;\n"


def to_excel(schema=SCHEMA, kind="FileOutputExcel", **config):
    """file -> Excel output, which writes out.xlsx."""
    sink = {"id": "out", "type": kind, "config": {"filename": "out.xlsx", **config},
            "schema": {"input": columns(schema), "output": []}, "inputs": ["row1"], "outputs": []}
    return job([reader(schema, header_rows=1), sink], [flow("row1", "in", "out")])


def picture(data):
    """What a workbook holds: each sheet's cells (value, type, number format) by place, and its column widths."""
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    seen = {"calculation": workbook.calculation.calcMode}
    for sheet in workbook.worksheets:
        cells = {cell.coordinate: (cell.value, cell.data_type, cell.number_format)
                 for row in sheet.iter_rows() for cell in row if cell.value is not None}
        widths = {letter: dimension.width for letter, dimension in sheet.column_dimensions.items() if dimension.width}
        seen[sheet.title] = {"cells": cells, "widths": widths, "rows": sheet.max_row, "columns": sheet.max_column}
    return seen


def on_both(tmp_path, made, data=DATA, files=None, name="out.xlsx"):
    """Both engines write the same workbook. Returns its picture, or None when neither wrote one."""
    inputs = {"in.csv": data, **(files or {})}
    key = answer_key.run_job(made, inputs, tmp_path / "v1", answer_key.run_v1)
    run = answer_key.run_job(made, inputs, tmp_path / "v2", answer_key.run_v2)
    assert key.succeeded and run.succeeded, (key.error, run.error)
    assert sorted(run.files) == sorted(key.files)
    if name not in run.files:
        return None
    assert picture(run.files[name]) == picture(key.files[name])
    return picture(run.files[name])


def values(sheet):
    return {place: value for place, (value, _, _) in sheet["cells"].items()}


def v2(tmp_path, made, data=DATA, files=None):
    """Run on v2 alone, inside tmp_path; returns how it ended and the picture of out.xlsx, None when there is none."""
    (tmp_path / "in.csv").write_bytes(data)
    for name, held in (files or {}).items():
        (tmp_path / name).write_bytes(held)
    previous = os.getcwd()
    os.chdir(tmp_path)
    try:
        result = run_job(made)
    finally:
        os.chdir(previous)
    written = tmp_path / "out.xlsx"
    return result, picture(written.read_bytes()) if written.exists() else None


def refusals(made):
    with pytest.raises(JobRefusedError) as refused:
        load_job(made)
    return [(refusal.key, refusal.reason) for refusal in refused.value.report]


# ------------------------------------------------------------------
# Rows, sheet and header
# ------------------------------------------------------------------

def test_rows_go_to_sheet1_with_no_header(tmp_path):
    seen = on_both(tmp_path, to_excel())
    assert list(seen) == ["calculation", "Sheet1"]
    assert values(seen["Sheet1"]) == {"A1": 1, "B1": "ann", "C1": 1.5, "A2": 2, "B2": "bob", "C2": 2, "A3": 3}


def test_header_is_the_names_of_the_columns(tmp_path):
    seen = on_both(tmp_path, to_excel(includeheader=True))
    assert [seen["Sheet1"]["cells"][place][0] for place in ("A1", "B1", "C1", "A2")] == ["id", "name", "amt", 1]


@pytest.mark.parametrize("given, name", [
    ("Data", "Data"), ("'Quoted'", "Quoted"), ('"Two words"', "Two words"), ("Sheet", "Sheet"), ("'Mixed\"", "'Mixed\""),
])
def test_sheet_is_named_and_quotes_around_the_name_are_no_part_of_it(tmp_path, given, name):
    assert list(on_both(tmp_path, to_excel(sheetname=given))) == ["calculation", name]


def test_first_cell_moves_everything_right_and_down(tmp_path):
    seen = on_both(tmp_path, to_excel(includeheader=True, first_cell_x="2", first_cell_y="3"))
    cells = values(seen["Sheet1"])
    assert cells["C4"] == "id" and cells["C5"] == 1 and cells["E5"] == 1.5 and "A1" not in cells
    assert values(on_both(tmp_path / "b", to_excel(first_cell_x=1, first_cell_y=None))["Sheet1"])["B1"] == 1
    assert values(on_both(tmp_path / "c", to_excel(first_cell_x="x", first_cell_y=""))["Sheet1"])["A1"] == 1


# ------------------------------------------------------------------
# What a cell holds
# ------------------------------------------------------------------

TYPED = "id:int, name:str, amt:float, m:Decimal#2, ok:bool, day:datetime@%d/%m/%Y, at:datetime, n:int"
TYPED_DATA = (b"id;name;amt;m;ok;day;at;n\n"
              b"1;ann;1.5;2.5;true;31/01/2024;2024-01-31 10:11:12;7\n"
              b"2;;;;;;;\n"
              b"3;cy;2;3;false;01/02/2024;2024-02-01;8\n")


def test_every_type_of_column(tmp_path):
    seen = on_both(tmp_path, to_excel(TYPED), TYPED_DATA)["Sheet1"]["cells"]
    assert seen["A1"] == (1, "n", "General") and seen["B1"] == ("ann", "s", "General")
    assert seen["D1"] == (2.5, "n", "0.00")
    assert seen["E1"] == (True, "b", "General") and seen["E3"] == (False, "b", "General")
    # A date with a pattern is text written by it, as in v1; one with none is a date of Excel's.
    assert seen["F1"] == ("31/01/2024", "s", "General")
    assert seen["G1"] == (dt.datetime(2024, 1, 31, 10, 11, 12), "d", "yyyy-mm-dd h:mm:ss")
    assert seen["G3"][0] == dt.datetime(2024, 2, 1)
    assert "B2" not in seen and "C2" not in seen and "G2" not in seen


def test_decimal_is_the_float_nearest_to_it(tmp_path):
    made = to_excel("a:Decimal#20, b:Decimal")
    data = b"a;b\n0.30000000000000004441;1234567890.123456789\n0.1;0.1\n7802576312314467;38624302505.203052\n"
    seen = on_both(tmp_path, made, data)["Sheet1"]["cells"]
    assert seen["A1"][0] == 0.3 and seen["B1"][0] == 1234567890.123457 and seen["A2"][0] == 0.1
    # Polars' own cast is a step away from the nearest float for these two.
    assert seen["A3"][0] == 7802576312314467 and seen["B3"][0] == 38624302505.20305


def test_number_format_is_of_the_declared_places(tmp_path):
    made = to_excel("a:Decimal#4, b:Decimal#0, c:Decimal, d:float#3, e:float, f:int#2")
    seen = on_both(tmp_path, made, b"a;b;c;d;e;f\n1.5;2;3.25;4.5;5.5;6\n")["Sheet1"]["cells"]
    assert [seen[place] for place in ("A1", "B1", "C1", "D1", "E1", "F1")] == [
        (1.5, "n", "0.0000"), (2, "n", "0"), (3.25, "n", "General"), (4.5, "n", "0.000"), (5.5, "n", "General"),
        (6, "n", "General")]


def test_columns_are_written_in_the_order_they_are_declared(tmp_path):
    made = to_excel(includeheader=True)
    made["components"][1]["schema"]["input"] = columns("amt:float, id:int, name:str")
    seen = on_both(tmp_path, made)
    assert [seen["Sheet1"]["cells"][place][0] for place in ("A1", "B1", "C1", "A2")] == ["amt", "id", "name", 1.5]


def test_row_with_nothing_in_it_is_written_as_in_talend(tmp_path):
    # v1 leaves such a row out and counts it as rejected.
    made = to_excel("a:str, b:int")
    made["components"][0]["config"]["remove_empty_row"] = False
    result, seen = v2(tmp_path, made, b"a;b\nx;1\n;\n  ;\ny;2\n")
    assert result.status == "success", result.error
    assert values(seen["Sheet1"]) == {"A1": "x", "B1": 1, "A3": "  ", "A4": "y", "B4": 2}


# ------------------------------------------------------------------
# A workbook that is there already
# ------------------------------------------------------------------

THERE = book({"Sheet1": [["id", "name", "amt"], [9, "old", 0.5]], "Other": [["kept"]]})


def test_workbook_that_is_there_is_replaced(tmp_path):
    seen = on_both(tmp_path, to_excel(), files={"out.xlsx": THERE})
    assert list(seen) == ["calculation", "Sheet1"] and values(seen["Sheet1"])["A1"] == 1


def test_append_file_adds_the_rows_under_those_that_are_there_and_no_second_header(tmp_path):
    seen = on_both(tmp_path, to_excel(includeheader=True, append_file=True), files={"out.xlsx": THERE})
    assert list(seen) == ["calculation", "Sheet1", "Other"]
    cells = values(seen["Sheet1"])
    assert cells["B2"] == "old" and cells["A3"] == 1 and cells["B5"] is None if "B5" in cells else True
    assert seen["Sheet1"]["rows"] == 5 and values(seen["Other"]) == {"A1": "kept"}


def test_append_file_into_a_sheet_that_is_not_there_makes_it_with_its_header(tmp_path):
    made = to_excel(includeheader=True, append_file=True, sheetname="New")
    seen = on_both(tmp_path, made, files={"out.xlsx": THERE})
    assert list(seen) == ["calculation", "Sheet1", "Other", "New"]
    assert values(seen["New"])["A1"] == "id" and values(seen["Sheet1"])["B2"] == "old"


def test_append_file_with_no_workbook_there_writes_a_new_one(tmp_path):
    seen = on_both(tmp_path, to_excel(includeheader=True, append_file=True))
    assert values(seen["Sheet1"])["A1"] == "id" and seen["Sheet1"]["rows"] == 4


def test_append_sheet_alone_does_not_keep_the_workbook(tmp_path):
    seen = on_both(tmp_path, to_excel(append_sheet=True), files={"out.xlsx": THERE})
    assert list(seen) == ["calculation", "Sheet1"] and values(seen["Sheet1"])["A1"] == 1


def test_appended_rows_start_under_the_last_row_and_at_the_first_cell_when_that_is_lower(tmp_path):
    made = to_excel(append_file=True, first_cell_y="6", first_cell_x="1")
    cells = values(on_both(tmp_path, made, files={"out.xlsx": THERE})["Sheet1"])
    assert cells["B7"] == 1 and cells["B8"] == 2
    made = to_excel(append_file=True, first_cell_y="0")
    assert values(on_both(tmp_path / "b", made, files={"out.xlsx": THERE})["Sheet1"])["A3"] == 1


# ------------------------------------------------------------------
# Widths, folders, no rows
# ------------------------------------------------------------------

FULL_DATA = (b"id;name;amt;m;ok;day;at;n\n"
             b"1;ann;1.5;2.5;true;31/01/2024;2024-01-31 10:11:12;7\n"
             b"30;christopher;2;13.255;false;01/02/2024;2024-02-01;0\n")


def test_every_column_is_as_wide_as_its_longest_value_when_asked(tmp_path):
    made = to_excel(TYPED, includeheader=True, is_all_auto_szie=True)
    seen = on_both(tmp_path, made, FULL_DATA)["Sheet1"]
    assert seen["widths"] == {"A": 4, "B": 13, "C": 5, "D": 7, "E": 6, "F": 12, "G": 21, "H": 3}
    made = to_excel(TYPED, is_all_auto_szie=True)
    assert on_both(tmp_path / "b", made, FULL_DATA)["Sheet1"]["widths"]["E"] == 6


def test_name_of_a_column_counts_for_its_width_only_when_every_column_is_sized(tmp_path):
    data = b"identifier;n\n1;abcdef\n"
    made = to_excel("identifier:int, n:str", includeheader=True, is_all_auto_szie=True)
    assert on_both(tmp_path, made, data)["Sheet1"]["widths"] == {"A": 12, "B": 8}
    made = to_excel("identifier:int, n:str", includeheader=True, auto_szie_setting=["identifier"])
    assert on_both(tmp_path / "b", made, data)["Sheet1"]["widths"] == {"A": 3}
    made = to_excel("identifier:int, n:str", is_all_auto_szie=True)
    assert on_both(tmp_path / "c", made, data)["Sheet1"]["widths"] == {"A": 3, "B": 8}


def test_only_the_columns_named_are_sized(tmp_path):
    made = to_excel(TYPED, includeheader=True, auto_szie_setting=["name", "m", "nope"], first_cell_x="1")
    seen = on_both(tmp_path, made, FULL_DATA)["Sheet1"]
    assert seen["widths"] == {"C": 13, "E": 7}


def test_missing_values_take_no_width(tmp_path):
    # v1 fails on a missing value when it sizes the columns ("boolean value of NA is ambiguous").
    result, seen = v2(tmp_path, to_excel(TYPED, is_all_auto_szie=True), TYPED_DATA)
    assert result.status == "success", result.error
    assert seen["Sheet1"]["widths"] == {"A": 3, "B": 5, "C": 5, "D": 6, "E": 6, "F": 12, "G": 21, "H": 3}


def test_folder_that_is_not_there_is_made(tmp_path):
    made = to_excel(filename="deep/er/out.xlsx")
    assert on_both(tmp_path, made, name="deep/er/out.xlsx") is not None


def test_folder_that_is_not_there_fails_the_job_when_it_is_not_to_be_made(tmp_path):
    result, seen = v2(tmp_path, to_excel(filename="deep/out.xlsx", create=False))
    assert result.status == "failed" and result.error == "Failed to write file 'deep/out.xlsx': folder 'deep' does not exist"


def test_no_rows_write_a_workbook_with_its_header_or_none_when_asked(tmp_path):
    seen = on_both(tmp_path, to_excel(includeheader=True), b"id;name;amt\n")
    assert values(seen["Sheet1"]) == {"A1": "id", "B1": "name", "C1": "amt"}
    assert on_both(tmp_path / "b", to_excel(includeheader=True, delete_empty_file=True), b"id;name;amt\n") is None
    # A workbook the run would have replaced is gone as well.
    made = to_excel(includeheader=True, delete_empty_file=True)
    on_both(tmp_path / "c", made, b"id;name;amt\n", files={"out.xlsx": THERE})
    assert not (tmp_path / "c" / "v1" / "out.xlsx").exists() and not (tmp_path / "c" / "v2" / "out.xlsx").exists()


def test_workbook_nothing_was_added_to_is_kept_as_in_talend(tmp_path):
    # v1 deletes the workbook it appended nothing to.
    made = to_excel(append_file=True, delete_empty_file=True)
    result, seen = v2(tmp_path, made, b"id;name;amt\n", files={"out.xlsx": THERE})
    assert result.status == "success" and values(seen["Sheet1"])["B2"] == "old" and list(seen)[1:] == ["Sheet1", "Other"]


def test_two_outputs_of_one_subjob_fill_two_sheets_of_one_workbook(tmp_path):
    # The second adds to the workbook the first has just written, as in v1, where they run one after the other.
    made = to_excel(sheetname="One", includeheader=True)
    second = dict(made["components"][1], id="out2", inputs=["row2"],
                  config={"filename": "out.xlsx", "sheetname": "Two", "append_file": True})
    made["components"][0]["outputs"] = ["row1", "row2"]
    made["components"].append(second)
    made["flows"].append(flow("row2", "in", "out2"))
    seen = on_both(tmp_path, made, files={"out.xlsx": THERE})
    assert list(seen) == ["calculation", "One", "Two"]
    assert values(seen["One"])["A1"] == "id" and values(seen["Two"])["A1"] == 1


def test_two_outputs_of_one_subjob_add_to_one_sheet(tmp_path):
    made = to_excel(includeheader=True, append_file=True)
    second = dict(made["components"][1], id="out2", inputs=["row2"])
    made["components"][0]["outputs"] = ["row1", "row2"]
    made["components"].append(second)
    made["flows"].append(flow("row2", "in", "out2"))
    seen = on_both(tmp_path, made)
    assert seen["Sheet1"]["rows"] == 7 and [values(seen["Sheet1"])[place] for place in ("A1", "A2", "A5")] == ["id", 1, 1]


def test_column_that_is_declared_and_does_not_arrive_is_empty_and_one_not_declared_is_not_written(tmp_path):
    made = to_excel(includeheader=True)
    made["components"][1]["schema"]["input"] = columns("name:str, nope:str, id:int")
    seen = on_both(tmp_path, made)
    assert values(seen["Sheet1"]) == {"A1": "name", "B1": "nope", "C1": "id", "A2": "ann", "C2": 1, "A3": "bob",
                                      "C3": 2, "C4": 3}


def test_columns_as_they_arrive_when_none_are_declared(tmp_path):
    made = to_excel(includeheader=True)
    made["components"][1]["schema"]["input"] = []
    seen = on_both(tmp_path, made)
    assert [values(seen["Sheet1"])[place] for place in ("A1", "B1", "C1", "C2")] == ["id", "name", "amt", 1.5]


def test_day_with_no_time_and_a_float_that_is_no_number(tmp_path):
    source = {"id": "in", "type": "PythonDataFrameComponent", "inputs": ["row0"], "outputs": ["row1"],
              "config": {"python_code": "df['f'] = float('nan')"},
              "schema": {"input": columns("id:int, d:date"), "output": columns("id:int, d:date, f:float")}}
    made = to_excel("id:int, d:date, f:float")
    made["components"][0] = reader("id:int, d:date", header_rows=1, outputs=("row0",))
    made["components"].insert(1, source)
    made["flows"] = [flow("row0", "in", "in2"), flow("row1", "in2", "out")]
    source["id"] = "in2"
    result, seen = v2(tmp_path, made, b"id;d\n1;2024-01-31\n")
    assert result.status == "success", result.error
    cells = seen["Sheet1"]["cells"]
    assert cells["B1"] == (dt.datetime(2024, 1, 31), "d", "yyyy-mm-dd h:mm:ss") and "C1" not in cells


def test_number_that_is_none_is_no_cell(tmp_path):
    # The logarithm of nothing has no end: a workbook cannot hold that, and a cell with no value in it is a fault.
    step = {"id": "it", "type": "Map", "inputs": ["row1"], "outputs": ["row2"],
            "config": {"inputs": {"main": {"name": "row1"}}, "outputs": [{"name": "row2", "columns": [
                {"name": "id", "expression": "row1.id"}, {"name": "f", "expression": "math.log(row1.amt - 1.5)"}]}]},
            "schema": {"inputs": {"row1": columns(SCHEMA)}}}
    made = to_excel()
    made["components"][1].update(inputs=["row2"], schema={"input": columns("id:int, f:float"), "output": []})
    made["components"].insert(1, step)
    made["flows"] = [flow("row1", "in", "it"), flow("row2", "it", "out")]
    result, seen = v2(tmp_path, made, b"id;name;amt\n1;ann;1.5\n2;bob;4.5\n")
    assert result.status == "success", result.error
    assert list(values(seen["Sheet1"])) == ["A1", "A2", "B2"]
    with zipfile.ZipFile(tmp_path / "out.xlsx") as held:
        assert "<v></v>" not in held.read("xl/worksheets/sheet1.xml").decode()


def test_file_that_is_no_workbook_cannot_be_added_to(tmp_path):
    (tmp_path / "in.csv").write_bytes(DATA)
    (tmp_path / "out.xlsx").write_bytes(b"not a workbook")
    previous = os.getcwd()
    os.chdir(tmp_path)
    try:
        result = run_job(to_excel(append_file=True))
    finally:
        os.chdir(previous)
    assert result.status == "failed"
    assert result.error == "'out.xlsx' is there and is no workbook that can be read: File is not a zip file"
    assert sorted(path.name for path in tmp_path.iterdir()) == ["in.csv", "out.xlsx"]
    assert (tmp_path / "out.xlsx").read_bytes() == b"not a workbook"


def test_sheet_too_long_for_a_workbook_fails_in_plain_words(tmp_path):
    result, seen = v2(tmp_path, to_excel(first_cell_y=1048574))
    assert result.status == "failed" and seen is None
    assert result.error == "a sheet holds 1048576 rows at most, and there are 3 to write"
    (tmp_path / "b").mkdir()
    (tmp_path / "b" / "in.csv").write_bytes(DATA)
    previous = os.getcwd()
    os.chdir(tmp_path / "b")
    try:
        result = run_job(to_excel(first_cell_y=1048573))
    finally:
        os.chdir(previous)
    # Read as a stream: a sheet this long is slow to take in whole.
    last = None
    for row in openpyxl.load_workbook(tmp_path / "b" / "out.xlsx", read_only=True)["Sheet1"].iter_rows(min_row=1048576):
        last = row[0].value
    assert result.status == "success" and last == 3
    result, seen = v2(tmp_path, to_excel(first_cell_y=1048573, includeheader=True))
    assert result.status == "failed" and result.error == "a sheet holds 1048576 rows at most, and there are 3 to write"


def test_formulas_are_worked_out_again_when_asked(tmp_path):
    assert on_both(tmp_path, to_excel(recalculate_formula=True))["calculation"] == "auto"


def test_failure_further_on_in_the_subjob_leaves_the_workbook_as_it_was(tmp_path):
    made = to_excel()
    made["components"][0]["config"]["die_on_error"] = True
    result, seen = v2(tmp_path, made, b"id;name;amt\n1;ann;1.5\nx;bob;2\n", files={"out.xlsx": THERE})
    assert result.status == "failed" and values(seen["Sheet1"])["B2"] == "old"
    assert sorted(path.name for path in tmp_path.iterdir()) == ["in.csv", "out.xlsx"]


def test_rows_written_are_in_globalmap_and_counted(tmp_path):
    made = dict(to_excel(includeheader=True), run={"row_counts": True})
    result, _ = v2(tmp_path, made)
    assert result.status == "success" and result.rows == {"out": 3}
    assert result.counts["out"] == {"NB_LINE": 3, "NB_LINE_OK": 3, "NB_LINE_REJECT": 0}


# ------------------------------------------------------------------
# What is refused, and what fails
# ------------------------------------------------------------------

def test_character_excel_cannot_hold_fails_in_plain_words(tmp_path):
    source = {"id": "in", "type": "FixedFlowInput", "inputs": [], "outputs": ["row1"],
              "config": {"use_inlinecontent": True, "inline_content": "1;ann|2;b\x01b|3;cy", "row_separator": "|"},
              "schema": {"input": [], "output": columns("id:int, name:str")}}
    made = to_excel("id:int, name:str")
    made["components"][0] = source
    result, seen = v2(tmp_path, made)
    assert result.status == "failed" and seen is None
    assert result.error == "column 'name' of row 2 holds a character a workbook cannot hold: 'b\\x01b'"


@pytest.mark.parametrize("config, key, said", [
    ({"sheetname": "a/b"}, "sheetname", "'a/b' is no name for a sheet: it holds '/'"),
    ({"sheetname": "x" * 32}, "sheetname", "a sheet's name has 31 characters at most, and this has 32"),
    ({"protect_file": True, "password": "x"}, "protect_file", None),
    ({"usestream": True}, "usestream", None),
    ({"filename": ""}, "path", "the workbook to write has no name"),
    ({"first_cell_x": "-1"}, "first_cell_x", "-1 is before the first column"),
    ({"nope": 1}, "nope", None),
])
def test_what_cannot_be_done_refuses_the_job(config, key, said):
    found = refusals(to_excel(**config))
    assert [path for path, _ in found] == [key]
    assert said is None or found[0][1] == said


def test_keys_v1_never_reads_change_nothing(tmp_path):
    unread = dict(version_2007=True, usestream=False, streamname="outputStream", first_cell_y_absolute=True,
                  keep_cell_formating=True, font="ARIAL", protect_file=False, password="", flushonrow=True,
                  flushonrow_num="100", advanced_separator=True, thousands_separator=".", decimal_separator=",",
                  truncate_exceeding_characters=True, encoding="ISO-8859-15", streaming_append=True,
                  use_shared_strings_table=True, die_on_error=False)
    assert on_both(tmp_path, to_excel(**unread)) == on_both(tmp_path / "b", to_excel())


def test_both_v1_names_and_v2s_own(tmp_path):
    for at, kind in enumerate(("file_output_excel", "FileOutputExcel", "tFileOutputExcel")):
        (tmp_path / str(at)).mkdir()
        result, seen = v2(tmp_path / str(at), to_excel(kind=kind))
        assert result.status == "success" and values(seen["Sheet1"])["A1"] == 1
